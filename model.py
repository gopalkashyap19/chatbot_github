import shutil
import os
import re
import threading
import time
import gc
import json

import torch

from pypdf import PdfReader
import docx

from transformers import (
    AutoTokenizer,
    AutoModelForCausalLM,
    TextIteratorStreamer,
    StoppingCriteria,
    StoppingCriteriaList
)

from langchain_text_splitters import RecursiveCharacterTextSplitter

from langchain_huggingface import HuggingFaceEmbeddings

from langchain_community.vectorstores import FAISS


# ============================================================
# CONFIGURATION
# ============================================================

EMBEDDING_MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"

MODEL_NAME = "Qwen/Qwen2.5-1.5B-Instruct"

BASE_DIR = os.path.dirname(
    os.path.abspath(__file__)
)


# ============================================================
# VECTOR STORE DIRECTORY
# ============================================================

VECTOR_STORE_DIR = os.path.join(
    BASE_DIR,
    "vector_store"
)


# ============================================================
# ACTIVE USER GUIDE INFORMATION
# ============================================================

ACTIVE_GUIDE_FILE = os.path.join(
    BASE_DIR,
    "active_user_guide.json"
)


# ============================================================
# UPLOAD DIRECTORY
# ============================================================

UPLOAD_FOLDER = os.path.join(
    BASE_DIR,
    "uploads"
)


# ============================================================
# GLOBAL VECTOR STORE
# ============================================================

vector_store = None


# ============================================================
# VECTOR STORE LOCK
# ============================================================

vector_store_lock = threading.Lock()


# ============================================================
# PER-USER GENERATION CONTROL
# ============================================================

generation_events = {}

generation_events_lock = threading.Lock()


def start_user_generation(user_id):

    """
    Create a stop-control event for a specific user.
    """

    with generation_events_lock:

        event = threading.Event()

        generation_events[str(user_id)] = event

        return event


def get_generation_event(user_id):

    """
    Get the generation stop event for a user.
    """

    with generation_events_lock:

        return generation_events.get(
            str(user_id)
        )


def stop_generation(user_id):

    """
    Stop generation only for the specified user.
    """

    with generation_events_lock:

        event = generation_events.get(
            str(user_id)
        )

        if event is not None:

            event.set()

            print(
                f"[GENERATION STOPPED] User: {user_id}"
            )

            return True

    return False


def remove_generation_event(user_id):

    """
    Remove generation event after generation finishes.
    """

    with generation_events_lock:

        generation_events.pop(
            str(user_id),
            None
        )


# ============================================================
# USER STOPPING CRITERIA
# ============================================================

class UserStoppingCriteria(StoppingCriteria):

    def __init__(self, user_id):

        self.user_id = str(user_id)


    def __call__(
        self,
        input_ids,
        scores,
        **kwargs
    ):

        event = get_generation_event(
            self.user_id
        )

        if event is not None and event.is_set():

            return True

        return False


# ============================================================
# LOAD EMBEDDING MODEL
# ============================================================

print(
    "[MODEL] Loading embedding model..."
)


embeddings = HuggingFaceEmbeddings(

    model_name=EMBEDDING_MODEL_NAME,

    model_kwargs={

        "device": "cpu"

    },

    encode_kwargs={

        "normalize_embeddings": True

    }

)


print(
    "[MODEL] Embedding model loaded."
)


# ============================================================
# LOAD QWEN MODEL
# ============================================================

print(
    "[MODEL] Loading Qwen generation model..."
)


tokenizer = AutoTokenizer.from_pretrained(
    MODEL_NAME
)


if tokenizer.pad_token is None:

    tokenizer.pad_token = tokenizer.eos_token


model = AutoModelForCausalLM.from_pretrained(

    MODEL_NAME,

    torch_dtype=torch.float32,

    device_map="cpu"

)


model.eval()


model.generation_config.max_length = None


print(
    "[MODEL] Qwen generation model loaded."
)


# ============================================================
# MODEL GENERATION LOCK
# ============================================================

model_generation_lock = threading.Lock()


# ============================================================
# LOAD EXISTING VECTOR STORE
# ============================================================

def load_existing_vector_store():

    """
    Load the FAISS vector database from disk.
    """

    global vector_store


    if not os.path.isdir(
        VECTOR_STORE_DIR
    ):

        print(
            "[GUIDE] No existing vector store found."
        )

        return None


    try:

        vector_store = FAISS.load_local(

            VECTOR_STORE_DIR,

            embeddings,

            allow_dangerous_deserialization=True

        )


        print(
            "[GUIDE] Existing vector database loaded successfully."
        )


        return vector_store


    except Exception as error:

        print(
            "[GUIDE] Could not load vector database:",
            error
        )


        vector_store = None


        return None


# ============================================================
# LOAD VECTOR STORE ON APPLICATION START
# ============================================================

load_existing_vector_store()


# ============================================================
# PDF TEXT EXTRACTION
# ============================================================

def extract_text_from_pdf(filepath):

    reader = PdfReader(
        filepath
    )


    text_parts = []


    for page in reader.pages:

        page_text = page.extract_text()


        if page_text:

            text_parts.append(
                page_text
            )


    return "\n".join(
        text_parts
    )


# ============================================================
# DOCX TEXT EXTRACTION
# ============================================================

def extract_text_from_docx(filepath):

    document = docx.Document(
        filepath
    )


    text_parts = []


    # --------------------------------------------------------
    # EXTRACT PARAGRAPHS
    # --------------------------------------------------------

    for paragraph in document.paragraphs:

        if paragraph.text.strip():

            text_parts.append(
                paragraph.text.strip()
            )


    # --------------------------------------------------------
    # EXTRACT TABLES
    # --------------------------------------------------------

    for table in document.tables:

        for row in table.rows:

            for cell in row.cells:

                if cell.text.strip():

                    text_parts.append(
                        cell.text.strip()
                    )


    return "\n".join(
        text_parts
    )


# ============================================================
# UNIVERSAL FILE TEXT EXTRACTION
# ============================================================

def extract_text_from_file(filepath):

    extension = os.path.splitext(
        filepath
    )[1].lower()


    if extension == ".pdf":

        return extract_text_from_pdf(
            filepath
        )


    elif extension == ".docx":

        return extract_text_from_docx(
            filepath
        )


    elif extension == ".txt":

        with open(

            filepath,

            "r",

            encoding="utf-8",

            errors="ignore"

        ) as file_handle:

            return file_handle.read()


    else:

        raise ValueError(

            f"Unsupported file type: {extension}"

        )


# ============================================================
# CREATE TEXT CHUNKS
# ============================================================

def create_documents_from_text(
    user_guide
):

    user_guide = re.sub(

        r"\s+",

        " ",

        user_guide

    ).strip()


    text_splitter = RecursiveCharacterTextSplitter(

        chunk_size=700,

        chunk_overlap=150,

        separators=[

            "\n\n",

            "\n",

            ". ",

            " ",

            ""

        ]

    )


    documents = text_splitter.create_documents(

        [user_guide]

    )


    return documents


# ============================================================
# PROCESS USER GUIDE
# ============================================================

def process_user_guide(
    user_guide,
    append=False
):

    """
    Process a company user guide.

    append=False:
        Replace the complete existing knowledge base.

    append=True:
        Add the new guide to the existing knowledge base.
    """

    global vector_store


    # ========================================================
    # VALIDATE USER GUIDE
    # ========================================================

    if not user_guide or not user_guide.strip():

        return {

            "success": False,

            "message":
                "User guide cannot be empty."

        }


    try:


        print("\n================================================")

        print("[GUIDE] Processing user guide...")

        print(f"[GUIDE] Append mode: {append}")

        print("================================================")


        # ====================================================
        # CREATE DOCUMENT CHUNKS
        # ====================================================

        documents = create_documents_from_text(
            user_guide
        )


        if not documents:

            return {

                "success": False,

                "message":
                    "Could not create document chunks."

            }


        print(
            f"[GUIDE] New chunks created: {len(documents)}"
        )


        # ====================================================
        # LOCK VECTOR STORE
        # ====================================================

        with vector_store_lock:


            # =================================================
            # APPEND MODE
            # =================================================

            if append:


                print(
                    "[GUIDE] Adding guide to existing knowledge base..."
                )


                # ---------------------------------------------
                # CASE 1:
                # VECTOR STORE ALREADY LOADED IN MEMORY
                # ---------------------------------------------

                if vector_store is not None:


                    vector_store.add_documents(
                        documents
                    )


                    print(
                        "[GUIDE] Added documents to in-memory knowledge base."
                    )


                # ---------------------------------------------
                # CASE 2:
                # VECTOR STORE EXISTS ON DISK
                # ---------------------------------------------

                elif os.path.isdir(
                    VECTOR_STORE_DIR
                ):


                    print(
                        "[GUIDE] Loading existing knowledge base from disk..."
                    )


                    try:


                        vector_store = FAISS.load_local(

                            VECTOR_STORE_DIR,

                            embeddings,

                            allow_dangerous_deserialization=True

                        )


                        print(
                            "[GUIDE] Existing knowledge base loaded."
                        )


                        vector_store.add_documents(
                            documents
                        )


                        print(
                            "[GUIDE] New guide appended successfully."
                        )


                    except Exception as error:


                        print(
                            "[GUIDE] ERROR: Could not load existing knowledge base:",
                            error
                        )


                        return {

                            "success": False,

                            "message":
                                "Existing knowledge base could not be loaded. "
                                "Your existing guides were not modified."

                        }


                # ---------------------------------------------
                # CASE 3:
                # NO EXISTING KNOWLEDGE BASE
                # ---------------------------------------------

                else:


                    print(
                        "[GUIDE] No existing knowledge base found."
                    )


                    vector_store = FAISS.from_documents(

                        documents,

                        embeddings

                    )


                    print(
                        "[GUIDE] Created first knowledge base."
                    )


            # =================================================
            # REPLACE MODE
            # =================================================

            else:


                print(
                    "[GUIDE] Replace mode enabled."
                )


                vector_store = FAISS.from_documents(

                    documents,

                    embeddings

                )


                print(
                    "[GUIDE] Previous knowledge base replaced."
                )


            # =================================================
            # CREATE VECTOR STORE DIRECTORY
            # =================================================

            os.makedirs(

                VECTOR_STORE_DIR,

                exist_ok=True

            )


            # =================================================
            # SAVE FAISS VECTOR STORE
            # =================================================

            vector_store.save_local(

                VECTOR_STORE_DIR

            )


            print(
                "[GUIDE] Knowledge base saved successfully."
            )


        print(
            "[GUIDE] User guide processing completed successfully."
        )


        return {

            "success": True,

            "message":
                (
                    "User guide added successfully."
                    if append
                    else
                    "User guide uploaded successfully and previous knowledge base replaced."
                ),

            "total_chunks":
                len(documents),

            "append":
                append

        }


    except Exception as error:


        print(
            "[GUIDE] PROCESSING ERROR:",
            error
        )


        return {

            "success": False,

            "message":
                str(error)

        }
# ============================================================
# REMOVE COMPANY USER GUIDE
# ============================================================

def remove_user_guide():

    """
    Remove:

    1. FAISS vector database
    2. All uploaded user guides
    3. active_user_guide.json
    """

    global vector_store


    try:


        print(
            "\n[REMOVE GUIDE] Starting complete user guide removal..."
        )


        # ----------------------------------------------------
        # REMOVE VECTOR STORE FROM MEMORY
        # ----------------------------------------------------

        with vector_store_lock:


            vector_store = None


            gc.collect()


            print(
                "[REMOVE GUIDE] Vector store removed from memory."
            )


        # ----------------------------------------------------
        # DELETE VECTOR STORE DIRECTORY
        # ----------------------------------------------------

        if os.path.exists(
            VECTOR_STORE_DIR
        ):


            print(

                "[REMOVE GUIDE] Removing vector store directory..."

            )


            deleted = False


            for attempt in range(5):


                try:


                    shutil.rmtree(

                        VECTOR_STORE_DIR

                    )


                    deleted = True


                    print(

                        "[REMOVE GUIDE] Vector store directory deleted."

                    )


                    break


                except PermissionError as error:


                    print(

                        f"[REMOVE GUIDE] Attempt "
                        f"{attempt + 1}/5 failed:",

                        error

                    )


                    gc.collect()


                    time.sleep(1)


            if not deleted:


                raise Exception(

                    "Could not delete the vector store directory. "

                    "It may be locked by another process."

                )


        # ----------------------------------------------------
        # DELETE UPLOADED FILES
        # ----------------------------------------------------

        if os.path.exists(
            UPLOAD_FOLDER
        ):


            try:


                for filename in os.listdir(
                    UPLOAD_FOLDER
                ):


                    filepath = os.path.join(

                        UPLOAD_FOLDER,

                        filename

                    )


                    if os.path.isfile(
                        filepath
                    ):


                        os.remove(
                            filepath
                        )


                        print(

                            "[REMOVE GUIDE] Deleted uploaded file:",

                            filename

                        )


                print(
                    "[REMOVE GUIDE] All uploaded guide files deleted."
                )


            except Exception as error:


                print(

                    "[REMOVE GUIDE] Error deleting uploaded files:",

                    error

                )


        # ----------------------------------------------------
        # DELETE ACTIVE GUIDE JSON
        # ----------------------------------------------------

        if os.path.exists(
            ACTIVE_GUIDE_FILE
        ):


            os.remove(
                ACTIVE_GUIDE_FILE
            )


            print(
                "[REMOVE GUIDE] Active guide information deleted."
            )


        print(
            "[REMOVE GUIDE] User guides successfully removed.\n"
        )


        return {

            "success": True,

            "message":
                "All user guides removed successfully."

        }


    except Exception as error:


        print(
            "[REMOVE GUIDE] Error:",
            error
        )


        return {

            "success": False,

            "message":
                str(error)

        }


def remove_specific_user_guide(filename):

    global vector_store

    try:

        print(
            f"\n[REMOVE GUIDE] Removing guide: {filename}"
        )


        # ==========================================
        # CHECK ACTIVE GUIDE FILE
        # ==========================================

        if not os.path.exists(ACTIVE_GUIDE_FILE):

            return {

                "success": False,

                "message": "No user guides found."

            }


        # ==========================================
        # LOAD GUIDE DATA
        # ==========================================

        with open(

            ACTIVE_GUIDE_FILE,

            "r",

            encoding="utf-8"

        ) as file_handle:

            guide_data = json.load(
                file_handle
            )


        guides = guide_data.get(

            "guides",

            []

        )


        # ==========================================
        # FIND GUIDE
        # ==========================================

        guide_to_remove = None


        remaining_guides = []


        for guide in guides:

            if guide.get("filename") == filename:

                guide_to_remove = guide

            else:

                remaining_guides.append(
                    guide
                )


        if guide_to_remove is None:

            return {

                "success": False,

                "message": "User guide not found."

            }


        print(

            f"[REMOVE GUIDE] Found: {filename}"

        )


        # ==========================================
        # DELETE ORIGINAL FILE
        # ==========================================

        filepath = guide_to_remove.get(
            "filepath"
        )


        if (

            filepath

            and

            os.path.exists(filepath)

        ):

            os.remove(
                filepath
            )


            print(

                "[REMOVE GUIDE] Original file deleted."

            )


        # ==========================================
        # REMOVE OLD VECTOR STORE
        # ==========================================

        with vector_store_lock:


            vector_store = None


            if os.path.exists(
                VECTOR_STORE_DIR
            ):

                shutil.rmtree(
                    VECTOR_STORE_DIR
                )


                print(

                    "[REMOVE GUIDE] Old vector store deleted."

                )


            # ==========================================
            # REBUILD KNOWLEDGE BASE
            # USING REMAINING GUIDES
            # ==========================================

            if remaining_guides:


                print(

                    "[REMOVE GUIDE] Rebuilding knowledge base..."

                )


                all_documents = []


                for guide in remaining_guides:


                    remaining_filepath = guide.get(
                        "filepath"
                    )


                    if (

                        remaining_filepath

                        and

                        os.path.exists(
                            remaining_filepath
                        )

                    ):


                        text = extract_text_from_file(

                            remaining_filepath

                        )


                        if text and text.strip():


                            documents = (

                                create_documents_from_text(

                                    text

                                )

                            )


                            all_documents.extend(

                                documents

                            )


                # ==========================================
                # CREATE NEW FAISS DATABASE
                # ==========================================

                if all_documents:


                    vector_store = FAISS.from_documents(

                        all_documents,

                        embeddings

                    )


                    os.makedirs(

                        VECTOR_STORE_DIR,

                        exist_ok=True

                    )


                    vector_store.save_local(

                        VECTOR_STORE_DIR

                    )


                    print(

                        "[REMOVE GUIDE] Knowledge base rebuilt."

                    )


            else:


                vector_store = None


                print(

                    "[REMOVE GUIDE] No guides remaining."

                )


        # ==========================================
        # UPDATE JSON
        # ==========================================

        if remaining_guides:


            updated_data = {

                "guides": remaining_guides

            }


            with open(

                ACTIVE_GUIDE_FILE,

                "w",

                encoding="utf-8"

            ) as file_handle:


                json.dump(

                    updated_data,

                    file_handle,

                    indent=4

                )


        else:


            # No guides left

            if os.path.exists(
                ACTIVE_GUIDE_FILE
            ):

                os.remove(
                    ACTIVE_GUIDE_FILE
                )


        return {

            "success": True,

            "message":

                f"{filename} removed successfully."

        }


    except Exception as error:


        print(

            "[REMOVE GUIDE ERROR]:",

            error

        )


        return {

            "success": False,

            "message": str(error)

        }
# ============================================================
# RETRIEVE CONTEXT
# ============================================================


def retrieve_context(
    question,
    k=3
):

    global vector_store


    if vector_store is None:

        return None


    with vector_store_lock:


        if vector_store is None:

            return None


        documents = vector_store.similarity_search(

            question,

            k=k

        )


    if not documents:

        return None


    context = "\n\n".join(

        document.page_content

        for document in documents

    )


    return context


# ============================================================
# BUILD QWEN PROMPT
# ============================================================

def build_chat_prompt(
    question,
    context
):

    system_message = (

        "You are a helpful customer support assistant. "

        "Answer the customer's question using ONLY the "

        "information available in the company user guide "

        "context.\n\n"

        "Rules:\n"

        "- Understand the customer's question.\n"

        "- Give a natural and helpful answer.\n"

        "- Use only information from the context.\n"

        "- Do not invent information.\n"

        "- Do not use outside knowledge.\n"

        "- Do not return the complete context.\n"

        "- Reply with ONLY the answer.\n"

        "- Do not simulate additional conversations.\n"

        "- Do not add Customer or Assistant labels.\n"

        "- If the answer is unavailable, say exactly:\n"

        "\"I couldn't find this information in the company "

        "user guide.\""

    )


    user_message = (

        f"COMPANY USER GUIDE CONTEXT:\n"

        f"{context}\n\n"

        f"CUSTOMER QUESTION:\n"

        f"{question}"

    )


    chat_prompt = tokenizer.apply_chat_template(

        [

            {

                "role": "system",

                "content": system_message

            },

            {

                "role": "user",

                "content": user_message

            }

        ],

        tokenize=False,

        add_generation_prompt=True

    )


    return chat_prompt


# ============================================================
# CLEAN GENERATED ANSWER
# ============================================================

def clean_answer(answer):

    stop_markers = [

        "\nCustomer:",

        "\nUser:",

        "\nAssistant:",

        "\nBot:",

        "<|im_start|>",

        "\n[Note",

        "\nThank you for choosing"

    ]


    for marker in stop_markers:


        if marker in answer:


            answer = answer.split(
                marker
            )[0]


    return answer.strip()


# ============================================================
# STREAM GENERATION
# ============================================================

def generate_response_stream(

    question,

    user_id,

    stream_callback=None

):

    """
    Generate the chatbot response token-by-token.
    """

    global vector_store


    # --------------------------------------------------------
    # CHECK KNOWLEDGE BASE
    # --------------------------------------------------------

    if vector_store is None:


        return {

            "success": False,

            "stopped": False,

            "answer":

                "The company user guide has not been processed yet."

        }


    # --------------------------------------------------------
    # RETRIEVE CONTEXT
    # --------------------------------------------------------

    context = retrieve_context(

        question,

        k=3

    )


    if not context:


        return {

            "success": False,

            "stopped": False,

            "answer":

                "I couldn't find relevant information in the company user guide."

        }


    # --------------------------------------------------------
    # CREATE PROMPT
    # --------------------------------------------------------

    chat_prompt = build_chat_prompt(

        question,

        context

    )


    # --------------------------------------------------------
    # CREATE STOP EVENT
    # --------------------------------------------------------

    stop_event = start_user_generation(

        user_id

    )


    full_answer = ""


    try:


        # ----------------------------------------------------
        # TOKENIZE
        # ----------------------------------------------------

        inputs = tokenizer(

            chat_prompt,

            return_tensors="pt"

        )


        input_ids = inputs[
            "input_ids"
        ]


        attention_mask = inputs[
            "attention_mask"
        ]


        # ----------------------------------------------------
        # STREAMER
        # ----------------------------------------------------

        streamer = TextIteratorStreamer(

            tokenizer,

            skip_prompt=True,

            skip_special_tokens=True

        )


        # ----------------------------------------------------
        # STOPPING CRITERIA
        # ----------------------------------------------------

        stopping_criteria = StoppingCriteriaList(

            [

                UserStoppingCriteria(
                    user_id
                )

            ]

        )


        generation_kwargs = {


            "input_ids":
                input_ids,


            "attention_mask":
                attention_mask,


            "max_new_tokens":
                300,


            "do_sample":
                True,


            "temperature":
                0.3,


            "repetition_penalty":
                1.15,


            "no_repeat_ngram_size":
                3,


            "eos_token_id":
                tokenizer.eos_token_id,


            "pad_token_id":
                tokenizer.pad_token_id,


            "streamer":
                streamer,


            "stopping_criteria":
                stopping_criteria

        }


        # ----------------------------------------------------
        # MODEL GENERATION THREAD
        # ----------------------------------------------------

        def run_generation():


            with model_generation_lock:


                with torch.no_grad():


                    model.generate(

                        **generation_kwargs

                    )


        generation_thread = threading.Thread(

            target=run_generation,

            daemon=True

        )


        generation_thread.start()


        # ----------------------------------------------------
        # STREAM TOKENS
        # ----------------------------------------------------

        for text_chunk in streamer:


            if stop_event.is_set():

                break


            full_answer += text_chunk


            if stream_callback:


                stream_callback(

                    full_answer

                )


        generation_thread.join()


        # ----------------------------------------------------
        # GENERATION STOPPED
        # ----------------------------------------------------

        if stop_event.is_set():


            print(

                f"[GENERATION] Cancelled for user: {user_id}"

            )


            return {

                "success": False,

                "stopped": True,

                "answer":
                    clean_answer(full_answer)

            }


        # ----------------------------------------------------
        # CLEAN ANSWER
        # ----------------------------------------------------

        full_answer = clean_answer(

            full_answer

        )


        if not full_answer:


            full_answer = (

                "I couldn't generate a response."

            )


        return {

            "success": True,

            "stopped": False,

            "answer":
                full_answer

        }


    except Exception as error:


        print(

            "[GENERATION ERROR]:",

            error

        )


        return {

            "success": False,

            "stopped": False,

            "error":
                str(error),

            "answer":

                "An error occurred while generating the response."

        }


    finally:


        remove_generation_event(

            user_id

        )


# ============================================================
# NON-STREAMING COMPATIBILITY FUNCTION
# ============================================================

def generate_response(

    question,

    user_id=None

):

    """
    Normal non-streaming compatibility function.
    """

    if user_id is None:


        user_id = "default_user"


    return generate_response_stream(

        question,

        user_id,

        stream_callback=None

    )


# ============================================================
# DEBUG CONTEXT FUNCTION
# ============================================================

def get_relevant_context(
    question
):

    return retrieve_context(

        question,

        k=3

    )
